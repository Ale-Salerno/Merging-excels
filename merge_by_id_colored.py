from __future__ import annotations

import argparse
import json
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


MASTER_DEFAULT = "Kopie von translations_export_20.05.2026.xlsx"
SOURCE_DEFAULTS = [
    "Review_translations08.04.2026 (002) - MENTIONED TERMS TO BE KEPT IN EN.xlsx",
    "Review_translations08.04.2026 (002) - TO TRANSLATE FROM DE WITH EN REFERENCE.xlsx",
    "Review_translations08.04.2026 (002) - TO TRANSLATE FROM EN AS CORRECTION.xlsx",
    "Review_translations08.04.2026 (002) - TO TRANSLATE FROM EN WITH DE REFERENCE.xlsx",
]


@dataclass
class SourceReport:
    source_file: str
    matched_ids: int
    colored_updates_applied: int
    missing_ids_in_master: int
    duplicate_ids_detected: int
    duplicate_ids: list[str]
    matched_ids_different_row_number: int


def normalize_id(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return str(value)


def is_colored(cell: Any) -> bool:
    fill = cell.fill
    if fill is None or not fill.fill_type or fill.fill_type == "none":
        return False

    fg = fill.fgColor
    if fg is None:
        return True
    if fg.type == "rgb":
        rgb = (fg.rgb or "").upper()
        return rgb not in {"", "00000000", "00FFFFFF", "FFFFFFFF"}
    if fg.type == "indexed":
        return fg.index not in {None, 0, 64}
    if fg.type == "theme":
        return True
    return True


def header_map(sheet: Worksheet) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, sheet.max_column + 1):
        header = sheet.cell(1, col).value
        if header is None:
            continue
        key = str(header).strip()
        if key and key not in mapping:
            mapping[key] = col
    return mapping


def source_columns_by_header(sheet: Worksheet) -> dict[str, list[int]]:
    mapping: dict[str, list[int]] = {}
    for col in range(1, sheet.max_column + 1):
        header = sheet.cell(1, col).value
        if header is None:
            continue
        key = str(header).strip()
        if not key:
            continue
        mapping.setdefault(key, []).append(col)
    return mapping


def id_row_map(sheet: Worksheet, id_col: int) -> tuple[dict[str, int], Counter[str]]:
    rows: dict[str, int] = {}
    duplicates: Counter[str] = Counter()
    for row in range(2, sheet.max_row + 1):
        raw_id = sheet.cell(row, id_col).value
        normalized = normalize_id(raw_id)
        if normalized is None:
            continue
        if normalized in rows:
            duplicates[normalized] += 1
        else:
            rows[normalized] = row
    return rows, duplicates


def merge_sources(
    master_path: Path,
    source_paths: list[Path],
    output_path: Path,
    report_path: Path,
) -> None:
    master_wb = load_workbook(master_path)
    master_sheet = master_wb["Sheet1"]

    master_headers = header_map(master_sheet)
    if "key" not in master_headers:
        raise ValueError("Master Sheet1 is missing 'key' column.")

    master_id_col = master_headers["key"]
    master_rows, master_duplicates = id_row_map(master_sheet, master_id_col)
    if master_duplicates:
        raise ValueError(f"Duplicate IDs found in master: {sorted(master_duplicates)}")

    reports: list[SourceReport] = []
    total_updates = 0

    for source_path in source_paths:
        source_wb = load_workbook(source_path, data_only=False)
        source_sheet = source_wb["Sheet1"]
        source_headers = source_columns_by_header(source_sheet)
        source_id_cols = source_headers.get("key")
        if not source_id_cols:
            raise ValueError(f"Source Sheet1 missing 'key' column: {source_path.name}")
        source_id_col = source_id_cols[0]

        duplicate_counts: Counter[str] = Counter()
        missing_ids: set[str] = set()
        matched_ids: set[str] = set()
        matched_diff_row = 0
        updates_applied = 0

        shared_headers = [
            header
            for header in source_headers
            if header in master_headers and header != "key"
        ]

        first_seen_row: dict[str, int] = {}
        for row in range(2, source_sheet.max_row + 1):
            source_id = normalize_id(source_sheet.cell(row, source_id_col).value)
            if source_id is None:
                continue

            if source_id in first_seen_row:
                duplicate_counts[source_id] += 1
                continue
            first_seen_row[source_id] = row

            master_row = master_rows.get(source_id)
            if master_row is None:
                missing_ids.add(source_id)
                continue

            matched_ids.add(source_id)
            if master_row != row:
                matched_diff_row += 1

            for header in shared_headers:
                target_col = master_headers[header]
                for source_col in source_headers[header]:
                    source_cell = source_sheet.cell(row, source_col)
                    if not is_colored(source_cell):
                        continue
                    master_cell = master_sheet.cell(master_row, target_col)
                    if master_cell.value != source_cell.value:
                        master_cell.value = source_cell.value
                        updates_applied += 1

        total_updates += updates_applied
        reports.append(
            SourceReport(
                source_file=source_path.name,
                matched_ids=len(matched_ids),
                colored_updates_applied=updates_applied,
                missing_ids_in_master=len(missing_ids),
                duplicate_ids_detected=len(duplicate_counts),
                duplicate_ids=sorted(duplicate_counts),
                matched_ids_different_row_number=matched_diff_row,
            )
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    master_wb.save(output_path)

    report_payload = {
        "sheet": "Sheet1",
        "merge_key": "key",
        "merge_mode": "id-based",
        "source_cell_filter": "colored-cells-only",
        "total_colored_updates_applied": total_updates,
        "per_source": [report.__dict__ for report in reports],
    }
    report_path.write_text(json.dumps(report_payload, indent=2, ensure_ascii=False), encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Merge review workbooks into master by ID using colored cells only (Sheet1)."
    )
    parser.add_argument(
        "--master",
        type=Path,
        default=Path(MASTER_DEFAULT),
        help="Master workbook path.",
    )
    parser.add_argument(
        "--sources",
        nargs="*",
        type=Path,
        default=[Path(name) for name in SOURCE_DEFAULTS],
        help="Source workbook paths.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(MASTER_DEFAULT),
        help="Output workbook path.",
    )
    parser.add_argument(
        "--report",
        type=Path,
        default=Path("merge_report.json"),
        help="Path for merge report JSON.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    merge_sources(
        master_path=args.master,
        source_paths=args.sources,
        output_path=args.output,
        report_path=args.report,
    )
    print(f"Merged workbook written to: {args.output}")
    print(f"Merge report written to: {args.report}")


if __name__ == "__main__":
    main()
